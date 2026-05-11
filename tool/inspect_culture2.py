# -*- coding: utf-8 -*-
import openpyxl
import json

wb = openpyxl.load_workbook('assets/culture_facillity_2025.xlsx', read_only=False)
out = {}
for i in [2, 3, 8]:
    name = wb.sheetnames[i]
    ws = wb[name]
    out[name] = {'row4': [], 'row5': [], 'sample_row6': []}
    for c in range(1, min(20, ws.max_column + 1)):
        v4 = ws.cell(4, c).value
        v5 = ws.cell(5, c).value
        v6 = ws.cell(6, c).value
        out[name]['row4'].append((c, str(v4)[:30] if v4 else None))
        out[name]['row5'].append((c, str(v5)[:30] if v5 else None))
        out[name]['sample_row6'].append((c, str(v6)[:40] if v6 else None))

with open('tool/culture_inspect2.json', 'w', encoding='utf-8') as f:
    json.dump(out, f, ensure_ascii=False, indent=2)
print('wrote')
