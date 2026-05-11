# -*- coding: utf-8 -*-
import openpyxl
import json

wb = openpyxl.load_workbook('assets/culture_facillity_2025.xlsx', read_only=False)
out = []
for i, name in enumerate(wb.sheetnames):
    ws = wb[name]
    info = {'idx': i, 'name': name, 'rows': ws.max_row, 'cols': ws.max_column}
    # Find header row (first row with > 5 non-empty cells)
    header_row = None
    for r in range(1, min(10, ws.max_row + 1)):
        non_empty = sum(1 for c in range(1, ws.max_column + 1) if ws.cell(r, c).value)
        if non_empty > 5:
            header_row = r
            break
    info['header_row'] = header_row
    if header_row:
        headers = []
        for c in range(1, min(35, ws.max_column + 1)):
            v = ws.cell(header_row, c).value
            if v is not None:
                headers.append({'col': c, 'name': str(v)[:30]})
        info['headers'] = headers
        # Sample data row
        if ws.max_row > header_row:
            sample = {}
            for c in range(1, min(35, ws.max_column + 1)):
                v = ws.cell(header_row + 1, c).value
                if v is not None:
                    sample[f'c{c}'] = str(v)[:50]
            info['sample'] = sample
    out.append(info)

with open('tool/culture_inspect.json', 'w', encoding='utf-8') as f:
    json.dump(out, f, ensure_ascii=False, indent=2)
print('wrote tool/culture_inspect.json')
