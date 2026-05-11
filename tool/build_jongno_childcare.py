# -*- coding: utf-8 -*-
"""
교육부 유치원알리미 공시 엑셀에서 종로구 유치원 추출.
주소 + 좌표 + 운영시간 + 정원/현원 + 유형 모두 포함.

Run: python tool/build_jongno_childcare.py
"""
import json
from pathlib import Path

import openpyxl

XLSX_PATH = Path("assets/일반 현황_20261_서울특별시.xlsx")
OUT_PATH = Path("assets/jongno_childcare.json")


def to_int(v):
    if v is None:
        return 0
    try:
        return int(float(v))
    except (ValueError, TypeError):
        return 0


def to_float(v):
    if v is None:
        return None
    try:
        return float(v)
    except (ValueError, TypeError):
        return None


def normalize_opertime(raw):
    """엑셀의 운영시간 형식이 '0800시0800분~1800시1800분' 같이 깨져있어 정리"""
    if not raw:
        return None
    s = str(raw).strip()
    if not s:
        return None
    # 0800시0800분~1800시1800분 → 08시00분~18시00분
    import re
    # 시간 4자리 + '시' + 분 4자리 + '분' 패턴
    def fix(m):
        hhmm1, hhmm2 = m.group(1), m.group(2)
        if len(hhmm1) == 4 and hhmm1[:2] == hhmm2[:2]:
            return f'{hhmm1[:2]}시{hhmm1[2:]}분'
        return m.group(0)
    s = re.sub(r'(\d{4})시(\d{4})분', fix, s)
    return s


def map_type_code(establish):
    """설립유형 → typeCode (모델 호환)"""
    s = str(establish or '')
    if s.startswith('공립'):
        return '1'  # 국공립
    if s.startswith('사립'):
        return '4'  # 민간
    return ''


def main():
    wb = openpyxl.load_workbook(XLSX_PATH, read_only=False)
    ws = wb.active

    out = []
    for i in range(4, ws.max_row + 1):
        addr = ws.cell(i, 9).value or ''
        if '종로구' not in str(addr):
            continue
        name = (ws.cell(i, 3).value or '').strip()
        if not name:
            continue
        establish = (ws.cell(i, 4).value or '').strip()

        # 정원 = col19 인가정원 우선, 없으면 연령별 합
        capacity = to_int(ws.cell(i, 19).value)
        if capacity == 0:
            capacity = sum(to_int(ws.cell(i, c).value) for c in (20, 21, 22, 23, 24))

        # 현원 = 연령별 원아수 합
        current = sum(to_int(ws.cell(i, c).value) for c in (25, 26, 27, 28, 29))

        out.append({
            "id": f"kg_{name}",
            "name": name,
            "addr": str(addr).strip(),
            "lat": to_float(ws.cell(i, 30).value),
            "lng": to_float(ws.cell(i, 31).value),
            "typeCode": map_type_code(establish),
            "establish": establish,
            "operatingHours": normalize_opertime(ws.cell(i, 13).value),
            "capacity": capacity,
            "currentCount": current,
            "hasCctv": False,
            "staffCount": 0,
            "tel": str(ws.cell(i, 10).value or '').strip(),
            "homepage": str(ws.cell(i, 12).value or '').strip() or None,
        })

    out.sort(key=lambda x: x["name"])
    OUT_PATH.write_text(
        json.dumps({"data": out}, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )

    print(f"종로구 유치원: {len(out)}개")
    for x in out:
        print(f'  - {x["name"]} | {x["establish"]} | {x["operatingHours"]} | '
              f'정원{x["capacity"]}/현원{x["currentCount"]} | ({x["lat"]:.4f}, {x["lng"]:.4f})')
    print(f"\nWrote {OUT_PATH}")


if __name__ == "__main__":
    main()
